from datetime import timedelta
from decimal import Decimal
from io import BytesIO

from django.conf import settings
from django.contrib.auth import authenticate
from django.db.models import Count, DecimalField, F, Q, Sum, Value
from django.db.models.functions import Coalesce, TruncDate
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from django.utils import timezone
from rest_framework import status
from rest_framework.authtoken.models import Token
from rest_framework.pagination import PageNumberPagination
from rest_framework.permissions import AllowAny
from rest_framework.response import Response
from rest_framework.throttling import ScopedRateThrottle
from rest_framework.views import APIView

from . import paystack
from .emails import member_progress, send_payment_request, send_receipt, send_welcome
from .models import Member, Payment
from .serializers import (
    ManualPaymentSerializer, MemberSerializer, MemberUpdateSerializer, PaymentRequestSerializer, PaymentSerializer,
    PayStartSerializer, RegisterSerializer, normalize_phone,
)


class Pager(PageNumberPagination):
    page_size = 50
    page_size_query_param = "page_size"
    max_page_size = 200


def members_with_totals():
    paid = Coalesce(
        Sum("payments__amount", filter=Q(payments__status=Payment.SUCCESS)),
        Value(Decimal("0")),
        output_field=DecimalField(max_digits=12, decimal_places=2),
    )
    return Member.objects.annotate(paid=paid)


# --------------------------------------------------------------------------- public


class ConfigView(APIView):
    authentication_classes = []
    permission_classes = [AllowAny]

    def get(self, request):
        return Response({"mission_name": settings.MISSION_NAME, "target": settings.TARGET_AMOUNT})


class RegisterView(APIView):
    authentication_classes = []
    permission_classes = [AllowAny]
    throttle_classes = [ScopedRateThrottle]
    throttle_scope = "register"

    def post(self, request):
        serializer = RegisterSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        member = serializer.save()
        send_welcome(member)
        return Response({"name": member.name, "code": member.code}, status=status.HTTP_201_CREATED)


class PayStartView(APIView):
    """A member enters their phone and an amount; we send them to Paystack checkout."""

    authentication_classes = []
    permission_classes = [AllowAny]
    throttle_classes = [ScopedRateThrottle]
    throttle_scope = "pay"

    def post(self, request):
        serializer = PayStartSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        try:
            payment, url = paystack.start_payment(serializer.validated_data["phone"], serializer.validated_data["amount"])
        except paystack.PaystackError as exc:
            return Response({"detail": str(exc)}, status=status.HTTP_502_BAD_GATEWAY)
        return Response({"authorization_url": url, "reference": payment.reference})


class PayVerifyView(APIView):
    """Called by the page Paystack redirects back to."""

    authentication_classes = []
    permission_classes = [AllowAny]

    def get(self, request, reference):
        payment = get_object_or_404(Payment.objects.select_related("member"), reference=reference)
        if payment.status == Payment.PENDING:
            try:
                paystack.verify_payment(reference)
            except paystack.PaystackError as exc:
                return Response({"detail": str(exc)}, status=status.HTTP_502_BAD_GATEWAY)
            payment.refresh_from_db()
        progress = member_progress(payment.member)
        return Response({
            "status": payment.status,
            "amount": payment.amount,
            "receipt_no": payment.receipt_no,
            "name": payment.member.name,
            "paid": progress["paid"],
            "target": progress["target"],
            "balance": progress["balance"],
            "percent": progress["percent"],
        })


class ProgressView(APIView):
    """A member signs in with the email and phone number they registered with."""

    authentication_classes = []
    permission_classes = [AllowAny]
    throttle_classes = [ScopedRateThrottle]
    throttle_scope = "pay"

    def post(self, request):
        phone = normalize_phone(request.data.get("phone", ""))
        email = str(request.data.get("email", "")).strip()
        member = (
            Member.objects.filter(phone=phone, email__iexact=email, is_active=True).first()
            if phone and email else None
        )
        if not member:
            return Response(
                {"detail": "We couldn't find a member with that email and phone number. Use the ones you registered with."},
                status=status.HTTP_404_NOT_FOUND,
            )
        paystack.sync_pending(member.payments.all())
        progress = member_progress(member)
        payments = member.payments.filter(status=Payment.SUCCESS).order_by("-paid_at", "-id")
        return Response({
            "first_name": member.first_name,
            "name": member.name,
            "code": member.code,
            "email": member.email,
            "phone": member.phone,
            "duration": member.get_duration_display(),
            "registered": member.created_at,
            **{key: progress[key] for key in ("paid", "target", "balance", "percent", "completed")},
            "payments": [
                {"receipt_no": p.receipt_no, "amount": p.amount, "paid_at": p.paid_at, "status": p.status,
                 "method": "Cash / manual" if p.source == "manual" else (p.channel or "paystack").replace("_", " ")}
                for p in payments
            ],
        })


class PaystackWebhookView(APIView):
    authentication_classes = []
    permission_classes = [AllowAny]

    def post(self, request):
        if not paystack.valid_signature(request.body, request.headers.get("X-Paystack-Signature")):
            return Response({"detail": "Invalid signature"}, status=status.HTTP_401_UNAUTHORIZED)
        event = request.data
        if event.get("event") == "charge.success":
            paystack.apply_transaction(event.get("data") or {})
        return Response({"status": "ok"})


# --------------------------------------------------------------------------- auth


class LoginView(APIView):
    authentication_classes = []
    permission_classes = [AllowAny]
    throttle_classes = [ScopedRateThrottle]
    throttle_scope = "login"

    def post(self, request):
        user = authenticate(username=request.data.get("username", ""), password=request.data.get("password", ""))
        if not user or not user.is_staff:
            return Response({"detail": "Wrong username or password."}, status=status.HTTP_400_BAD_REQUEST)
        token, _ = Token.objects.get_or_create(user=user)
        return Response({"token": token.key, "username": user.get_username()})


class MeView(APIView):
    def get(self, request):
        return Response({"username": request.user.get_username()})


class LogoutView(APIView):
    def post(self, request):
        Token.objects.filter(user=request.user).delete()
        return Response(status=status.HTTP_204_NO_CONTENT)


# --------------------------------------------------------------------------- dashboard


class DashboardView(APIView):
    def get(self, request):
        paystack.sync_pending(Payment.objects.all())
        active = members_with_totals().filter(is_active=True)
        count = active.count()
        target_total = active.aggregate(t=Sum("target"))["t"] or Decimal("0")
        raised = (
            Payment.objects.filter(status=Payment.SUCCESS, member__is_active=True).aggregate(s=Sum("amount"))["s"]
            or Decimal("0")
        )
        completed = active.filter(paid__gte=F("target")).count()
        unpaid = active.filter(paid=0).count()
        by_duration = {row["duration"]: row["n"] for row in Member.objects.filter(is_active=True).values("duration").annotate(n=Count("id"))}

        today = timezone.localdate()
        start = today - timedelta(days=13)
        rows = (
            Payment.objects.filter(status=Payment.SUCCESS, paid_at__date__gte=start)
            .annotate(day=TruncDate("paid_at"))
            .values("day")
            .annotate(total=Sum("amount"))
        )
        totals = {row["day"]: row["total"] for row in rows}
        daily = [
            {"date": (start + timedelta(days=i)).isoformat(), "total": totals.get(start + timedelta(days=i), 0)}
            for i in range(14)
        ]
        recent = Payment.objects.filter(status=Payment.SUCCESS).select_related("member").order_by("-paid_at")[:8]
        return Response({
            "members": count,
            "target_total": target_total,
            "raised": raised,
            "percent": int(min(100, raised / target_total * 100)) if target_total else 0,
            "completed": completed,
            "unpaid": unpaid,
            "in_progress": count - completed - unpaid,
            "by_duration": {"1wk": by_duration.get("1wk", 0), "2wks": by_duration.get("2wks", 0)},
            "pending": Payment.objects.filter(status=Payment.PENDING).count(),
            "daily": daily,
            "recent": PaymentSerializer(recent, many=True).data,
        })


# --------------------------------------------------------------------------- members


class MemberListView(APIView):
    def get(self, request):
        qs = members_with_totals()
        params = request.query_params

        search = params.get("q", "").strip()
        if search:
            cond = (
                Q(name__icontains=search) | Q(email__icontains=search) | Q(code__icontains=search)
                | Q(phone__icontains=search)
            )
            phone = normalize_phone(search)
            if phone:
                cond |= Q(phone=phone)
            qs = qs.filter(cond)

        state = params.get("status")
        if state == "completed":
            qs = qs.filter(paid__gte=F("target"))
        elif state == "unpaid":
            qs = qs.filter(paid=0)
        elif state == "in_progress":
            qs = qs.filter(paid__gt=0, paid__lt=F("target"))

        if params.get("duration") in ("1wk", "2wks"):
            qs = qs.filter(duration=params["duration"])
        if params.get("active") != "all":
            qs = qs.filter(is_active=True)

        ordering = params.get("ordering", "-created_at")
        if ordering not in ("name", "-name", "paid", "-paid", "created_at", "-created_at"):
            ordering = "-created_at"
        qs = qs.order_by(ordering, "id")

        pager = Pager()
        page = pager.paginate_queryset(qs, request)
        return pager.get_paginated_response(MemberSerializer(page, many=True).data)


class MemberDetailView(APIView):
    def _get(self, pk):
        return get_object_or_404(members_with_totals(), pk=pk)

    def _payload(self, member):
        data = MemberSerializer(member).data
        data["payments"] = PaymentSerializer(member.payments.select_related("member"), many=True).data
        return data

    def get(self, request, pk):
        return Response(self._payload(self._get(pk)))

    def patch(self, request, pk):
        member = self._get(pk)
        serializer = MemberUpdateSerializer(member, data=request.data, partial=True)
        serializer.is_valid(raise_exception=True)
        serializer.save()
        return Response(self._payload(self._get(pk)))


# --------------------------------------------------------------------------- payments


class PaymentListView(APIView):
    def get(self, request):
        paystack.sync_pending(Payment.objects.all())
        qs = Payment.objects.select_related("member")
        params = request.query_params
        if params.get("status"):
            qs = qs.filter(status__in=params["status"].split(","))
        search = params.get("q", "").strip()
        if search:
            qs = qs.filter(
                Q(reference__icontains=search) | Q(receipt_no__icontains=search)
                | Q(member__name__icontains=search) | Q(member__phone__icontains=search)
            )
        pager = Pager()
        page = pager.paginate_queryset(qs, request)
        response = pager.get_paginated_response(PaymentSerializer(page, many=True).data)
        response.data["summary"] = payment_summary(qs)
        return response

    def post(self, request):
        """Treasurer records a payment made outside Paystack (for example cash)."""
        serializer = ManualPaymentSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        payment = serializer.save()
        send_receipt(payment)
        return Response(PaymentSerializer(payment).data, status=status.HTTP_201_CREATED)


def payment_summary(filtered):
    """Mission-wide paid / target / balance, plus the total of the payments in the current filter."""
    active = list(members_with_totals().filter(is_active=True).only("target"))
    target = sum((m.target for m in active), Decimal("0"))
    paid = sum((m.paid for m in active), Decimal("0"))
    balance = sum((max(m.target - m.paid, Decimal("0")) for m in active), Decimal("0"))
    return {
        "paid": paid,
        "target": target,
        "balance": balance,
        "shown_amount": filtered.aggregate(s=Sum("amount"))["s"] or Decimal("0"),
        "shown_count": filtered.count(),
    }


class PaymentDetailView(APIView):
    def get(self, request, pk):
        payment = get_object_or_404(Payment.objects.select_related("member"), pk=pk)
        if payment.status == Payment.PENDING and payment.source == "paystack":
            paystack.sync_pending(Payment.objects.filter(pk=pk))
            payment.refresh_from_db()
        data = PaymentSerializer(payment).data
        data["member"] = MemberSerializer(members_with_totals().get(pk=payment.member_id)).data
        return Response(data)


class PaymentRequestView(APIView):
    def post(self, request):
        serializer = PaymentRequestSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        member, amount = serializer.validated_data["member"], serializer.validated_data["amount"]
        message = serializer.validated_data.get("message", "").strip()
        note = f"Requested by {request.user.get_username()}"
        try:
            payment, url = paystack.start_payment(member, amount, note=note)
        except paystack.PaystackError as exc:
            return Response({"detail": f"Could not create the payment link: {exc}"}, status=status.HTTP_502_BAD_GATEWAY)
        send_payment_request(payment, url, message)
        result = {"emailed": member.email, "reference": payment.reference, "mpesa": None}
        if serializer.validated_data["mpesa"]:
            try:
                _, text = paystack.mpesa_prompt(member, amount, note=note + " (M-Pesa prompt)")
                result["mpesa"] = {"sent": True, "detail": text}
            except paystack.PaystackError as exc:
                result["mpesa"] = {"sent": False, "detail": str(exc)}
        return Response(result, status=status.HTTP_201_CREATED)


class PaymentReverseView(APIView):
    def post(self, request, pk):
        payment = get_object_or_404(Payment, pk=pk)
        if payment.status != Payment.SUCCESS:
            return Response({"detail": "Only paid payments can be reversed."}, status=status.HTTP_400_BAD_REQUEST)
        note = str(request.data.get("note", "")).strip()
        payment.status = Payment.REVERSED
        payment.note = (payment.note + "\n" if payment.note else "") + f"Reversed: {note or 'no reason given'}"
        payment.save()
        return Response(PaymentSerializer(payment).data)


class PaymentResendView(APIView):
    def post(self, request, pk):
        payment = get_object_or_404(Payment.objects.select_related("member"), pk=pk)
        if payment.status != Payment.SUCCESS:
            return Response({"detail": "Only paid payments have a receipt."}, status=status.HTTP_400_BAD_REQUEST)
        send_receipt(payment)
        return Response({"detail": "Receipt sent."})


# --------------------------------------------------------------------------- export


class ExportView(APIView):
    def get(self, request):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        wb = Workbook()
        members_sheet = wb.active
        members_sheet.title = "Members"
        members_sheet.append([
            "Code", "Name", "Phone", "Email", "Duration", "Target (KES)", "Paid (KES)",
            "Balance (KES)", "Percent", "Status", "Registered",
        ])
        for m in members_with_totals().filter(is_active=True).order_by("name"):
            data = MemberSerializer(m).data
            members_sheet.append([
                m.code, m.name, m.phone, m.email, m.get_duration_display(), float(m.target),
                float(data["paid"]), float(data["balance"]), data["percent"] / 100, data["status"].replace("_", " "),
                timezone.localtime(m.created_at).strftime("%Y-%m-%d %H:%M"),
            ])
        for row in members_sheet.iter_rows(min_row=2, min_col=9, max_col=9):
            for cell in row:
                cell.number_format = "0%"

        pay_sheet = wb.create_sheet("Payments")
        pay_sheet.append(["Receipt", "Reference", "Date", "Amount (KES)", "Member", "Phone", "Status", "Source", "Channel"])
        for p in Payment.objects.select_related("member"):
            pay_sheet.append([
                p.receipt_no, p.reference, timezone.localtime(p.paid_at or p.created_at).strftime("%Y-%m-%d %H:%M"),
                float(p.amount), p.member.name, p.member.phone, p.get_status_display(), p.get_source_display(), p.channel,
            ])

        header_fill = PatternFill("solid", fgColor="0E4568")
        for sheet in (members_sheet, pay_sheet):
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = header_fill
            for column in sheet.columns:
                width = max(len(str(c.value or "")) for c in column) + 2
                sheet.column_dimensions[column[0].column_letter].width = min(max(width, 10), 40)
            sheet.freeze_panes = "A2"

        buffer = BytesIO()
        wb.save(buffer)
        response = HttpResponse(
            buffer.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = f'attachment; filename="gongoni27-{timezone.localdate():%Y%m%d}.xlsx"'
        return response
